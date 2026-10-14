"""Convert the IEEE LaTeX paper to a readable .docx.

The PDF remains the submission artefact; this export exists so the paper can
be opened, read and commented on in Word. Citations are resolved to the same
numbers bibtex assigned, so [7] in Word is [7] in the PDF.
"""
import io
import os
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

PAPER = r"g:\Uni\AgriQ An AI and Quantum-Optimized System for Smart Crop Planning\agriq\paper"
TEX = os.path.join(PAPER, "main.tex")
BBL = os.path.join(PAPER, "main.bbl")
OUT = os.path.join(PAPER, "main.docx")
B = chr(92)

tex = io.open(TEX, encoding="utf-8").read()


def expand_inputs(src, depth=0):
    r"""Inline \input{...} so sections kept in their own file are exported.

    The Methodology lives in methodology.tex and is pulled in by main.tex at
    build time; without this the Word copy silently lost an entire section,
    one figure and one table while still looking well-formed.
    """
    if depth > 5:
        return src

    def one(m):
        name = m.group(1).strip()
        path = os.path.join(PAPER, name if name.endswith(".tex") else name + ".tex")
        if not os.path.exists(path):
            print(f"  WARNING: \\input{{{name}}} not found, section will be missing")
            return ""
        print(f"  inlined {os.path.basename(path)}")
        return expand_inputs(io.open(path, encoding="utf-8").read(), depth + 1)

    return re.sub(r"\\input\{([^}]+)\}", one, src)


tex = expand_inputs(tex)

# ---------------------------------------------------------------- citations
bbl = io.open(BBL, encoding="utf-8").read()
keys = re.findall(r"\\bibitem\{([^}]+)\}", bbl)
CITE = {k: i + 1 for i, k in enumerate(keys)}

# reference list text, one entry per bibitem
entries = re.split(r"\\bibitem\{[^}]+\}", bbl)[1:]


def clean_ref(t):
    t = re.sub(r"\\urlprefix|\\newblock", " ", t)
    t = re.sub(r"\\url\{([^}]*)\}", r"\1", t)
    t = re.sub(r"\\BIBentry[A-Za-z]*", "", t)
    t = re.sub(r"\\[A-Za-z]+\s*", "", t)
    t = t.replace("{", "").replace("}", "").replace("~", " ")
    t = t.replace("\\", "")
    t = t.replace("``", "“").replace("''", "”")
    t = t.replace("---", "—").replace("--", "–")
    return re.sub(r"\s+", " ", t).strip(" .,") + "."


REFS = [clean_ref(e) for e in entries]

# ---------------------------------------------------------------- inline tex
GREEK = {"lambda": "\u03bb", "beta": "\u03b2", "phi": "\u03c6", "ell": "\u2113",
         "pm": "\u00b1", "times": "\u00d7", "circ": "\u00b0", "sigma": "\u03c3",
         "mathrm": "", "emph": "", "textbf": "", "text": ""}


def inline(t):
    """LaTeX fragment -> readable unicode text."""
    t = re.sub(r"\\cite\{([^}]*)\}",
               lambda m: "[" + ", ".join(
                   str(CITE.get(k.strip(), "?")) for k in m.group(1).split(",")) + "]", t)
    t = re.sub(r"\\ref\{(?:tab|fig|sec):([^}]*)\}", lambda m: m.group(1), t)
    t = t.replace(r"\refpending", "[ref. pending]")
    t = re.sub(r"\\(?:emph|textbf|textit)\{([^}]*)\}", r"\1", t)
    t = re.sub(r"\\mathrm\{([^}]*)\}", r"\1", t)
    t = re.sub(r"\$([^$]*)\$", r"\1", t)
    for k, v in GREEK.items():
        t = t.replace("\\" + k, v)
    t = re.sub(r"\^\{?([0-9-]+)\}?", lambda m: m.group(1)
               .translate(str.maketrans("0123456789-", "\u2070\u00b9\u00b2\u00b3\u2074"
                                        "\u2075\u2076\u2077\u2078\u2079\u207b")), t)
    t = re.sub(r"_\{?([A-Za-z0-9]+)\}?", r"\1", t)
    t = t.replace("---", "\u2014").replace("--", "\u2013")
    t = t.replace("``", "\u201c").replace("''", "\u201d")
    t = t.replace(r"\%", "%").replace(r"\$", "$").replace(r"\&", "&")
    t = t.replace(r"\,", " ").replace("~", " ").replace(r"\ ", " ")
    t = t.replace("{,}", ",").replace("{", "").replace("}", "")
    t = re.sub(r"\\[A-Za-z]+", "", t)
    # LaTeX hard-wraps source lines; Word paragraphs must reflow, so collapse
    # every run of whitespace (newlines included) to a single space.
    return re.sub(r"\s+", " ", t).strip()


# ---------------------------------------------------------------- document
doc = Document()
st = doc.styles["Normal"]
st.font.name = "Times New Roman"
st.font.size = Pt(10)

sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Inches(0.8)
sec.left_margin = sec.right_margin = Inches(0.9)


def para(text, size=10, bold=False, italic=False, align=None, space_after=6):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p


title = inline(re.search(r"\\title\{(.*?)\n*\}", tex, re.S).group(1).replace(r"\\", " "))
para(title, size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
def byline():
    r"""Read the author block that main.tex's \ifblind toggle selects.

    This was hardcoded to "Anonymous Submission", so the Word copy carried the
    blinded byline even when the PDF was the named one. Read the toggle and
    take the matching branch instead.
    """
    m = re.search(r"\\blind(true|false)\s*%\s*AUTHORSHIP-TOGGLE", tex)
    blind = (m.group(1) == "true") if m else True
    branch = re.search(
        r"\\ifblind(.*?)\\else(.*?)\\fi" if not blind else r"\\ifblind(.*?)\\else",
        tex, re.S)
    if not branch:
        return "Anonymous Submission  |  NGIA-QC 2026 \u2014 Double-blind review"
    block = branch.group(1) if blind else branch.group(2)
    parts = re.findall(r"\\IEEEauthorblock[NA]\{(.*?)\}", block, re.S)
    return "  |  ".join(inline(p) for p in parts if inline(p))


para(byline(), size=10, italic=True,
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)

abstract = inline(re.search(r"\\begin\{abstract\}(.*?)\\end\{abstract\}", tex, re.S).group(1))
p = doc.add_paragraph()
r = p.add_run("Abstract\u2014")
r.bold = True
r.font.size = Pt(9)
r = p.add_run(abstract)
r.font.size = Pt(9)
p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p.paragraph_format.space_after = Pt(6)

kw = inline(re.search(r"\\begin\{IEEEkeywords\}(.*?)\\end\{IEEEkeywords\}", tex, re.S).group(1))
p = doc.add_paragraph()
r = p.add_run("Index Terms\u2014")
r.bold = True
r.font.size = Pt(9)
r = p.add_run(kw)
r.font.size = Pt(9)
r.italic = True
p.paragraph_format.space_after = Pt(14)

# ---------------------------------------------------------------- tables
ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII"]


def add_table(caption, body, label):
    n = ROMAN[add_table.count]
    add_table.count += 1
    para("TABLE %s. %s" % (n, inline(caption)), size=9, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=3)
    rows = []
    for line in body.split(r"\\"):
        line = line.strip()
        if not line or line.startswith("%"):
            continue
        for cmd in (r"\toprule", r"\midrule", r"\bottomrule", r"\hline"):
            line = line.replace(cmd, "")
        line = line.strip()
        if not line:
            continue
        cells = []
        for c in line.split("&"):
            mc = re.search(r"\\multicolumn\{(\d+)\}\{[^}]*\}\{([^}]*)\}", c)
            if mc:
                # a multicolumn spans N columns; emit the text then N-1 blanks
                cells.append(inline(mc.group(2)))
                cells.extend([""] * (int(mc.group(1)) - 1))
            else:
                cells.append(inline(c))
        if any(c for c in cells):
            rows.append(cells)
    if not rows:
        return
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j in range(ncol):
            cell = t.cell(i, j)
            cell.text = row[j] if j < len(row) else ""
            for pp in cell.paragraphs:
                for rr in pp.runs:
                    rr.font.size = Pt(8)
                    rr.bold = (i == 0)
    doc.add_paragraph().paragraph_format.space_after = Pt(8)


add_table.count = 0

FIGNO = {"fig_map.png": 1, "fig_encoding.png": 2, "fig_water.png": 3}


def add_figure(path, caption):
    f = os.path.basename(path)
    doc.add_picture(os.path.join(PAPER, "figures", f), width=Inches(4.4))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    para("Fig. %d. %s" % (FIGNO.get(f, 0), inline(caption)), size=9,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)


# ---------------------------------------------------------------- body walk
body = tex.split(r"\section{Introduction}", 1)[1]
body = r"\section{Introduction}" + body.split(r"\bibliographystyle")[0]

token = re.compile(
    r"(\\section\{[^}]*\}|\\subsection\{[^}]*\}"
    r"|\\begin\{table\*?\}.*?\\end\{table\*?\}"
    r"|\\begin\{figure\}.*?\\end\{figure\}"
    r"|\\begin\{equation\}.*?\\end\{equation\})", re.S)

for chunk in token.split(body):
    chunk = chunk.strip()
    if not chunk:
        continue
    if chunk.startswith(r"\section{"):
        para(inline(chunk[9:-1]).upper(), size=11, bold=True, space_after=6)
    elif chunk.startswith(r"\subsection{"):
        para(inline(chunk[12:-1]), size=10, bold=True, italic=True, space_after=5)
    elif chunk.startswith(r"\begin{table"):
        cap = re.search(r"\\caption\{(.*?)\}\s*\n", chunk, re.S)
        lab = re.search(r"\\label\{([^}]*)\}", chunk)
        tab = re.search(r"\\begin\{tabular\}\{[^}]*\}(.*?)\\end\{tabular\}", chunk, re.S)
        if cap and tab:
            add_table(cap.group(1), tab.group(1), lab.group(1) if lab else "")
    elif chunk.startswith(r"\begin{figure}"):
        img = re.search(r"\\includegraphics\[[^\]]*\]\{figures/([^}]*)\}", chunk)
        cap = re.search(r"\\caption\{(.*?)\}\s*\\label", chunk, re.S)
        if img and cap:
            add_figure(img.group(1).replace(".pdf", ".png"), cap.group(1))
    elif chunk.startswith(r"\begin{equation}"):
        eq = inline(chunk[16:-14])
        para(eq, size=10, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
    else:
        for block in re.split(r"\n\s*\n", chunk):
            txt = inline(block)
            if len(txt) > 2:
                p = doc.add_paragraph()
                r = p.add_run(txt)
                r.font.size = Pt(10)
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.first_line_indent = Inches(0.2)

para("REFERENCES", size=11, bold=True, space_after=6)
for i, ref in enumerate(REFS, 1):
    p = doc.add_paragraph()
    r = p.add_run("[%d] %s" % (i, ref))
    r.font.size = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.first_line_indent = Inches(-0.3)

note = doc.add_paragraph()
r = note.add_run(
    "\nNote: this Word file is a readable export of main.tex for review and "
    "comment. The submission artefact is main.pdf, which is the authoritative "
    "IEEE two-column version.")
r.font.size = Pt(8)
r.italic = True
r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

doc.save(OUT)
print("wrote", OUT)
print("  tables:", add_table.count, " references:", len(REFS))
